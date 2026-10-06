"""Builds the mid-term and end-term presentations on the UPES project template.

    uv run --with python-pptx python tools/build_slides.py

Numbers are read from results/mac/*/*.json, so re-running after more experiments finish updates
every chart and table. Output: presentation/KAN_Rec_{MidTerm,EndTerm}.pptx
"""
import json
import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
PRES = ROOT / "presentation"
TEMPLATE = PRES / "templates" / "Project Presentation Template.pptx"
ASSETS = PRES / "assets"
RESULTS = ROOT / "results" / "mac"
REPO = "https://github.com/yashdagar/kan-rec"

TITLE = "Kolmogorov-Arnold Networks in Recurrent Recommenders: Session-Based and Next-Basket Prediction"
STUDENTS = [("Abhishek Yadav", "R2142231410", "500124297"),
            ("Saksham Agrawal", "R2142231383", "500121969"),
            ("Yash Dagar", "R2142231424", "500125147")]

BLUE = RGBColor(0x46, 0xB0, 0xFA)
NAVY = RGBColor(0x1B, 0x2A, 0x41)
TEXT = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x6B, 0x72, 0x80)
KAN = RGBColor(0xE8, 0x6A, 0x10)
MLP = RGBColor(0x1E, 0x7F, 0xC8)
BASE = RGBColor(0x8A, 0x94, 0xA3)
GOOD = RGBColor(0x2E, 0x8B, 0x57)
TINT = RGBColor(0xEE, 0xF6, 0xFE)
KANTINT = RGBColor(0xFD, 0xF0, 0xE6)
GREY_TINT = RGBColor(0xF3, 0xF4, 0xF6)
PURPLE = RGBColor(0x8E, 0x24, 0xAA)
ORANGE = RGBColor(0xF2, 0x8C, 0x28)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"


def load(dataset, stem):
    p = RESULTS / dataset / f"{stem}.json"
    return json.loads(p.read_text()) if p.exists() else None


def metric(dataset, stem, key, split="test"):
    r = load(dataset, stem)
    return None if r is None else r[split][key]



def _font(run, size, bold=False, color=TEXT, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def _bullet(paragraph, char="•", indent=0.22, level=0):
    pPr = paragraph._p.get_or_add_pPr()
    left = int(Inches(indent * (level + 1)))
    pPr.set("marL", str(left))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", FONT)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def text(slide, x, y, w, h, paras, size=16, color=TEXT, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT,
         space_after=6, name=None):
    """paras: list of str or dict(text|runs, size, bold, color, bullet, level, italic, align, space)."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, p in enumerate(paras):
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", align)
        para.space_after = Pt(p.get("space", space_after))
        runs = p.get("runs") or [(p.get("text", ""), {})]
        for rt, ro in runs:
            r = para.add_run()
            r.text = rt
            _font(r, ro.get("size", p.get("size", size)), ro.get("bold", p.get("bold", False)),
                  ro.get("color", p.get("color", color)), ro.get("italic", p.get("italic", False)))
        if p.get("bullet"):
            _bullet(para, level=p.get("level", 0), char="–" if p.get("level", 0) else "•")
    return tb


def box(slide, x, y, w, h, fill=TINT, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08, name=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        s.name = name
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def label_box(slide, x, y, w, h, lines, fill=TINT, line=None, size=14, color=TEXT, bold_first=True,
              align=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = box(slide, x, y, w, h, fill=fill, line=line, shape=shape)
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.08))
    for side in ("margin_top", "margin_bottom"):
        setattr(tf, side, Inches(0.04))
    for i, ln in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run()
        r.text = ln
        _font(r, size if i == 0 else size - 2, bold=bold_first and i == 0, color=color)
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("h", "med")
    return c


def circle_num(slide, x, y, n, fill=BLUE, d=0.42, size=14):
    s = box(slide, x, y, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    tf = s.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    _font(r, size, bold=True, color=WHITE)
    return s


def table(slide, x, y, w, rows, col_w=None, size=12, header_fill=NAVY, row_h=0.36, bold_rows=(),
          highlight=None, align_right_from=1):
    nr, nc = len(rows), len(rows[0])
    gs = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
    tbl = gs.table
    tblPr = tbl._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tblPr.set(attr, "0")
    style = tblPr.find(qn("a:tableStyleId"))
    if style is None:
        style = etree.SubElement(tblPr, qn("a:tableStyleId"))
    style.text = "{5940675A-B579-460E-94D1-54222C63F5DA}"
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
    for r in range(nr):
        tbl.rows[r].height = Inches(row_h)
        for c in range(nc):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = rows[r][c]
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if (c >= align_right_from and r > 0) else (
                PP_ALIGN.CENTER if (c >= align_right_from and r == 0) else PP_ALIGN.LEFT)
            run = p.add_run()
            run.text = "" if val is None else str(val)
            fill = None
            if r == 0:
                _font(run, size, bold=True, color=WHITE)
                fill = header_fill
            else:
                bold = r in bold_rows
                col = TEXT
                if highlight and (r, c) in highlight:
                    col, bold = highlight[(r, c)], True
                _font(run, size, bold=bold, color=col)
                fill = GREY_TINT if r % 2 == 0 else WHITE
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill
    return gs


def bar_chart(slide, x, y, w, h, cats, series, colors=None, point_colors=None, number_format="0.000",
              vmin=None, vmax=None, legend=False, horizontal=True, font_size=12, gap=60, title=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(font_size)
    ch.font.color.rgb = TEXT
    ch.has_legend = legend
    if legend:
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(font_size)
    if title:
        ch.has_title = True
        ch.chart_title.text_frame.text = title
        tp = ch.chart_title.text_frame.paragraphs[0]
        for r in tp.runs:
            _font(r, font_size + 1, bold=True, color=NAVY)
    else:
        ch.has_title = False
    plot = ch.plots[0]
    plot.gap_width = gap
    plot.overlap = -10 if len(series) > 1 else 0
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = number_format
    dl.number_format_is_linked = False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(font_size - 1)
    dl.font.color.rgb = TEXT
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
    va.format.line.fill.background()
    va.tick_labels.font.size = Pt(font_size - 2)
    va.tick_labels.font.color.rgb = MUTED
    va.tick_labels.number_format = number_format.replace("000", "00")
    va.tick_labels.number_format_is_linked = False
    if vmin is not None:
        va.minimum_scale = vmin
    if vmax is not None:
        va.maximum_scale = vmax
    ca = ch.category_axis
    ca.tick_labels.font.size = Pt(font_size)
    ca.format.line.color.rgb = RGBColor(0xC9, 0xCE, 0xD6)
    ca.has_major_gridlines = False
    if horizontal:
        ca.reverse_order = True
    for si, s in enumerate(plot.series):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = (colors or [MLP])[si % len(colors or [MLP])]
        if point_colors and len(series) == 1:
            for pi, col in enumerate(point_colors):
                pt = s.points[pi]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = col
    return gf


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s



def _set_runs(paragraph, txt):
    runs = paragraph.runs
    runs[0].text = txt
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def open_template():
    prs = Presentation(TEMPLATE)
    sld_ids = prs.slides._sldIdLst
    ids = list(sld_ids)
    keep = {0, len(ids) - 1}
    for i, sid in enumerate(ids):
        if i not in keep:
            prs.part.drop_rel(sid.get(qn("r:id")))
            sld_ids.remove(sid)
    from pptx.opc.packuri import PackURI
    prs.slides[1].part.partname = PackURI("/ppt/slides/slide2.xml")
    content_layout = [l for l in prs.slide_layouts if l.name == "Title and Content"][0]
    return prs, content_layout


def move_to_end(prs, slide):
    sld_ids = prs.slides._sldIdLst
    for sid in list(sld_ids):
        if prs.part.related_part(sid.get(qn("r:id"))) is slide.part:
            sld_ids.remove(sid)
            sld_ids.append(sid)
            return


def fill_title_slide(slide, stage, stage_color):
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text
        if t.startswith("Major/Minor"):
            _set_runs(sh.text_frame.paragraphs[0], "Major Project")
        elif t.startswith("Title"):
            sh.top = Inches(2.62)
            sh.left = Inches(1.0)
            sh.width = Inches(11.33)
            p = sh.text_frame.paragraphs[0]
            _set_runs(p, "Title: " + TITLE)
            p.alignment = PP_ALIGN.CENTER
            for r in p.runs:
                r.font.size = Pt(24)
        elif t.startswith("Presented"):
            paras = sh.text_frame.paragraphs
            sh.width = Inches(6.8)
            for i, (name, roll, sap) in enumerate(STUDENTS):
                _set_runs(paras[i + 1], f"{name}  |  {roll}  |  SAP {sap}")
            for p in paras[len(STUDENTS) + 1:]:
                p._p.getparent().remove(p._p)
        elif t.startswith("Mentored"):
            sh.text_frame.word_wrap = True
            sh.height = Inches(1.2)
            p0 = sh.text_frame.paragraphs[0]
            _set_runs(p0, "Mentored by:")
            ref = p0.runs[0]
            for line in ("Dr. Amar Jindal", "Assistant Professor,", "Data Science Cluster"):
                p = sh.text_frame.add_paragraph()
                r = p.add_run()
                r.text = line
                r.font.name = ref.font.name
                r.font.size = ref.font.size
    text(slide, 3.4, 3.75, 6.53, 0.5, [{"text": stage, "bold": True, "color": stage_color, "size": 20,
                                         "align": PP_ALIGN.CENTER}], name="Stage")


def new_slide(prs, layout, title):
    s = prs.slides.add_slide(layout)
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    text(s, 0.36, 0.26, 10.4, 0.62, [{"text": title, "bold": True, "color": BLUE, "size": 30}],
         anchor=MSO_ANCHOR.MIDDLE, name="Title")
    return s



def s_contents(prs, L, sections):
    s = new_slide(prs, L, "Content")
    n = len(sections)
    cols = 2 if n > 6 else 1
    per = math.ceil(n / cols)
    for i, sec in enumerate(sections):
        c, r = divmod(i, per)
        x, y = 1.2 + c * 5.6, 1.45 + r * 0.82
        circle_num(s, x, y, i + 1, fill=BLUE)
        text(s, x + 0.6, y - 0.02, 4.6, 0.46, [{"text": sec, "size": 20, "color": NAVY, "bold": True}],
             anchor=MSO_ANCHOR.MIDDLE)
    return s


def s_background(prs, L):
    s = new_slide(prs, L, "1. Introduction: Technical Background")
    cards = [
        ("Session-based recommendation", "Anonymous browsing session",
         ["Input: the clicks so far, i1, i2, ..., it", "Output: a ranking of all items for the next click",
          "No user profile: visitors are anonymous", "Model: GRU4Rec (gated recurrent unit)",
          "Data: YooChoose (RecSys Challenge 2015)"]),
        ("Next-basket recommendation", "Returning grocery customer",
         ["Input: the customer's past baskets B1 ... BT", "Output: a ranking of all items for basket BT+1",
          "Long histories, most items are re-purchases", "Model: basket-GRU (van Maasakkers et al., 2023)",
          "Data: Instacart, Dunnhumby Complete Journey"]),
    ]
    for i, (head, sub, items) in enumerate(cards):
        x = 0.7 + i * 6.1
        box(s, x, 1.3, 5.8, 4.25, fill=TINT)
        text(s, x + 0.3, 1.45, 5.2, 0.5, [{"text": head, "bold": True, "size": 20, "color": NAVY}])
        text(s, x + 0.3, 1.95, 5.2, 0.4, [{"text": sub, "italic": True, "size": 14, "color": MUTED}])
        text(s, x + 0.3, 2.45, 5.25, 3.0, [{"text": t, "bullet": True, "size": 16} for t in items],
             space_after=8)
    box(s, 0.7, 5.8, 11.9, 1.0, fill=KANTINT)
    text(s, 1.0, 5.85, 11.3, 0.9, [{"runs": [
        ("Common core: ", {"bold": True, "color": KAN}),
        ("both tasks use a GRU. This project asks where, if anywhere, Kolmogorov-Arnold Network (KAN) "
         "layers improve these GRU recommenders when compared fairly.", {})], "size": 16}],
         anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Two sequential recommendation tasks; both are modelled with GRUs, which gives a common place to insert KAN layers.")


def s_concepts(prs, L):
    s = new_slide(prs, L, "1. Introduction: Technical Concepts Used")
    text(s, 0.7, 1.25, 6.0, 4.2, [
        {"text": "Kolmogorov-Arnold Network (KAN) layer", "bold": True, "size": 18, "color": NAVY, "space": 4},
        {"text": "A learnable univariate function on every edge:", "size": 15, "space": 2},
        {"text": "y_o = Σ_i φ_o,i(x_i),   φ(x) = w·silu(x) + Σ_k c_k B_k(x)", "size": 15,
         "italic": True, "color": KAN, "space": 10},
        {"text": "Four basis families B_k implemented: cubic B-spline, Gaussian RBF (FastKAN), Chebyshev, "
                 "group-rational (GR-KAN)", "bullet": True, "size": 15},
        {"text": "GRU4Rec: official PyTorch code, session-parallel mini-batches, 2048 shared negatives, "
                 "cross-entropy with logQ correction", "bullet": True, "size": 15},
        {"text": "Basket-GRU: multi-hot basket in, sigmoid score for every product out", "bullet": True,
         "size": 15},
        {"text": "Parameter-matched MLP control for every KAN block, so gains from extra parameters are "
                 "not credited to the KAN", "bullet": True, "size": 15},
    ], space_after=7)
    s.shapes.add_picture(str(ASSETS / "kan_bases.png"), Inches(6.95), Inches(1.6), width=Inches(5.75))
    text(s, 6.95, 3.15, 5.75, 0.4, [{"text": "Basis functions used inside the KAN layers (from our code)",
                                       "size": 11, "color": MUTED, "align": PP_ALIGN.CENTER}])
    box(s, 6.95, 3.75, 5.75, 1.75, fill=GREY_TINT)
    text(s, 7.2, 3.85, 5.3, 1.6, [
        {"text": "Residual KAN block", "bold": True, "size": 15, "color": NAVY, "space": 4},
        {"text": "h' = h + KAN(LayerNorm(h))", "italic": True, "size": 15, "color": KAN, "space": 4},
        {"text": "Output layer initialised to zero, so each variant starts exactly as its baseline.",
         "size": 13},
    ])
    notes(s, "KAN: learnable functions on edges instead of fixed activations on nodes.")


def s_motivation(prs, L):
    s = new_slide(prs, L, "1. Introduction: Motivation and Problem Statement")
    text(s, 0.7, 1.3, 5.9, 0.45, [{"text": "Motivation", "bold": True, "size": 20, "color": NAVY}])
    items = [
        "Anonymous and privacy-restricted sessions make profile-based recommenders unusable",
        "KANs are claimed to be more accurate per parameter and more interpretable than MLPs",
        "Later studies (Yu et al. 2024; Hou et al. 2025) find KANs rarely win when parameters are matched",
        "Our first draft compared KAN and GRU on a leaky split with a broken B-spline, so its numbers could not "
        "support any conclusion",
    ]
    text(s, 0.7, 1.85, 5.9, 4.5, [{"text": t, "bullet": True, "size": 16} for t in items], space_after=10)
    box(s, 6.95, 1.3, 5.75, 5.2, fill=TINT)
    text(s, 7.25, 1.5, 5.15, 0.45, [{"text": "Problem Statement", "bold": True, "size": 20, "color": NAVY}])
    text(s, 7.25, 2.1, 5.15, 4.3, [
        {"text": "Determine whether, and where, KAN layers improve GRU-based session and next-basket "
                 "recommenders, under a protocol that allows a fair conclusion:", "size": 16, "space": 10},
        {"text": "standard, leakage-free splits comparable with published work", "bullet": True, "size": 15},
        {"text": "verified implementations of baselines and KAN layers", "bullet": True, "size": 15},
        {"text": "parameter-matched MLP controls for every KAN", "bullet": True, "size": 15},
        {"text": "evaluation on both session and basket data", "bullet": True, "size": 15},
    ], space_after=8)


def s_application_data(prs, L):
    s = new_slide(prs, L, "1. Introduction: Applications and Datasets")
    apps = [("E-commerce", "next product in a browsing session"),
            ("Online grocery", "pre-filled next basket, reminders"),
            ("Streaming and news", "next video, song or article"),
            ("Privacy-first platforms", "no long-term tracking needed")]
    text(s, 0.7, 1.25, 4.3, 0.45, [{"text": "Area of application", "bold": True, "size": 18, "color": NAVY}])
    for i, (a, b) in enumerate(apps):
        y = 1.8 + i * 1.05
        circle_num(s, 0.7, y + 0.08, i + 1, fill=ORANGE)
        text(s, 1.25, y, 3.8, 0.9, [{"text": a, "bold": True, "size": 15, "space": 2},
                                     {"text": b, "size": 13, "color": MUTED}])
    text(s, 5.3, 1.25, 7.3, 0.45, [{"text": "Datasets and input format", "bold": True, "size": 18,
                                     "color": NAVY}])
    rows = [["", "YooChoose 1/64", "Instacart", "Dunnhumby"],
            ["Task", "next click", "next basket", "next basket"],
            ["Raw input", "session, time, item", "user, order, product", "household, basket, day, product"],
            ["Model input", "item ID sequence", "multi-hot baskets", "multi-hot baskets"],
            ["Users / sessions", "139,628 sessions", "20,620 users*", "2,478 households"],
            ["Items", "17,264", "12,058 groups", "10,370"],
            ["Test targets", "55,357 clicks", "10,310 baskets", "1,239 baskets"],
            ["Split", "by time (last day)", "last basket, users 50/50", "last basket, users 50/50"]]
    table(s, 5.3, 1.8, 7.4, rows, col_w=[1.7, 1.9, 1.9, 1.9], size=12, row_h=0.5, align_right_from=9)
    text(s, 5.3, 5.95, 7.4, 0.6, [{"text": "* Mac profile: random 10% of Instacart users; the GPU profile "
                                            "uses all 206,209. Preprocessing follows GRU4Rec / SR-GNN and "
                                            "van Maasakkers et al. (2023).", "size": 11, "color": MUTED}])


def s_lit_table(prs, L):
    s = new_slide(prs, L, "2. Literature Review: Related Work")
    rows = [["Work", "Year", "Model", "Relevance / limitation"],
            ["Hidasi et al.", "2016/18", "GRU4Rec", "Strong session baseline; tuned version used here"],
            ["Hidasi & Czapp", "2023", "Reimplementation study", "Third-party GRU4Rec ports score 30-99% lower"],
            ["Wu et al. (SR-GNN)", "2019", "Graph NN", "Standard yoochoose1/64 benchmark table"],
            ["Hu et al. (TIFU-KNN)", "2020", "Frequency + kNN", "Simple model beats most neural NBR models"],
            ["Li et al.", "2023", "NBR reality check", "Repeat items dominate next-basket accuracy"],
            ["van Maasakkers et al.", "2023", "Basket-GRU", "Our basket baseline and protocol"],
            ["Liu et al. (KAN)", "2025", "KAN", "Learnable edge functions; slow to train"],
            ["Yu et al.", "2024", "KAN vs MLP", "With matched size, KAN wins only on symbolic formulas"],
            ["Genet & Inzirillo (TKAN)", "2024", "KAN in LSTM", "KAN inside a gated recurrent cell"],
            ["Park et al. (CF-KAN)", "2024", "KAN autoencoder", "Only KAN recommender with a matched MLP"],
            ["Xu et al. (KANRec)", "2026", "KAN interest fusion", "KAN in sequential recommendation"]]
    table(s, 0.7, 1.2, 11.95, rows, col_w=[2.75, 1.1, 2.5, 5.6], size=12, row_h=0.43, align_right_from=9)


def s_lit_inference(prs, L):
    s = new_slide(prs, L, "2. Literature Review: Inference and Research Gap")
    groups = [("Session-based models", BLUE, [
        "GRU4Rec remains a strong baseline when the official code is used",
        "Weak GRU4Rec results in many papers come from flawed reimplementations"]),
        ("Next-basket models", BLUE, [
            "Repeat purchases carry most of the signal",
            "Personal-frequency baselines and TIFU-KNN are very hard to beat"]),
        ("KANs", KAN, [
            "Interpretable learned functions, but slower than MLPs",
            "Gains often vanish against an MLP of the same size"])]
    for i, (h, col, items) in enumerate(groups):
        x = 0.7 + i * 4.05
        box(s, x, 1.3, 3.8, 2.75, fill=TINT if col == BLUE else KANTINT)
        text(s, x + 0.25, 1.45, 3.3, 0.45, [{"text": h, "bold": True, "size": 17, "color": NAVY}])
        text(s, x + 0.25, 2.0, 3.35, 2.0, [{"text": t, "bullet": True, "size": 14} for t in items],
             space_after=8)
    box(s, 0.7, 4.3, 11.95, 2.4, fill=GREY_TINT)
    text(s, 1.0, 4.42, 11.4, 2.2, [
        {"text": "Research gap", "bold": True, "size": 18, "color": NAVY, "space": 6},
        {"text": "No published work places KAN layers inside GRU4Rec or applies KANs to next-basket "
                 "recommendation", "bullet": True, "size": 15},
        {"text": "Few KAN recommenders control for model size with a parameter-matched MLP",
         "bullet": True, "size": 15},
        {"runs": [("Our question: ", {"bold": True, "color": KAN}),
                  ("where in a GRU recommender does a KAN layer help, compared with an MLP of the same size, on "
                   "standard benchmarks?", {})], "size": 15},
    ], space_after=6)


def s_swot(prs, L):
    s = new_slide(prs, L, "2. Literature Review: SWOT Analysis")
    quad = [("Strengths", GOOD, ["Official GRU4Rec and published protocols", "Parameter-matched MLP control for every KAN",
                                 "Unit-tested KAN layer with four bases", "Three datasets, two tasks"]),
            ("Weaknesses", KAN, ["KAN layers are slower to train", "Extra hyperparameters (basis, grid size)",
                                 "Mac runs limited to one seed and a 10% Instacart sample"]),
            ("Opportunities", MLP, ["Interpretable models for repeat purchasing (KAN-NBR)",
                                    "GPU runs on the UPES Blackwell cluster", "KAN in attention models (SASRec)"]),
            ("Threats", PURPLE, ["Results may show no KAN advantage over MLPs",
                                 "Transformer recommenders dominate recent work", "Compute time for many variants"])]
    for i, (h, col, items) in enumerate(quad):
        r, c = divmod(i, 2)
        x, y = 0.7 + c * 6.05, 1.25 + r * 2.85
        box(s, x, y, 5.85, 2.65, fill=GREY_TINT)
        circle_num(s, x + 0.25, y + 0.22, h[0], fill=col, d=0.5, size=18)
        text(s, x + 0.9, y + 0.24, 4.7, 0.46, [{"text": h, "bold": True, "size": 18, "color": NAVY}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.3, y + 0.9, 5.35, 1.7, [{"text": t, "bullet": True, "size": 14} for t in items],
             space_after=5)


def s_objectives(prs, L):
    s = new_slide(prs, L, "3. Objectives")
    box(s, 0.7, 1.25, 11.95, 1.3, fill=KANTINT)
    text(s, 1.0, 1.32, 11.4, 1.2, [
        {"text": "Main objective", "bold": True, "size": 17, "color": KAN, "space": 4},
        {"text": "Build and fairly evaluate KAN-augmented GRU recommenders for session-based and next-basket "
                 "prediction, and report where KAN layers help and where they do not.", "size": 16}])
    subs = ["Correct the methodological and reporting errors of the first draft",
            "Implement and unit-test a KAN layer with B-spline, RBF, Chebyshev and rational bases",
            "Add KAN layers to official GRU4Rec in three places: head, input, GRU cell",
            "Apply the same placements to the basket-GRU on Instacart and Dunnhumby",
            "Compare with parameter-matched MLPs and strong non-neural baselines, with significance tests",
            "Propose KAN-NBR, an interpretable additive KAN over repeat-purchase features"]
    text(s, 0.7, 2.8, 6, 0.45, [{"text": "Sub-objectives", "bold": True, "size": 17, "color": NAVY}])
    for i, t in enumerate(subs):
        r, c = divmod(i, 2)
        x, y = 0.7 + c * 6.05, 3.35 + r * 1.12
        circle_num(s, x, y + 0.1, i + 1, fill=BLUE)
        text(s, x + 0.6, y, 5.3, 0.95, [{"text": t, "size": 15}], anchor=MSO_ANCHOR.MIDDLE)


def s_reference_model(prs, L):
    s = new_slide(prs, L, "4. Methodology: Reference Software Model")
    text(s, 0.7, 1.2, 11.9, 0.45, [{"text": "KAN placements in GRU4Rec (same placements in the basket-GRU)",
                                     "bold": True, "size": 17, "color": NAVY}])
    labels = [("Item i_t", "(or basket)"), ("Embedding", "E[i_t]"), ("Input block", "S2: KAN / MLP"),
              ("GRU cell", "S3: KAN-GRU cell"), ("Head block", "S1: KAN / MLP"), ("Scores", "E·h' + b")]
    fills = [GREY_TINT, GREY_TINT, KANTINT, TINT, KANTINT, GREY_TINT]
    xs = [0.7, 2.75, 4.8, 7.05, 9.3, 11.05]
    widths = [1.65, 1.65, 1.85, 1.85, 1.35, 1.6]
    for i, ((a, b), f) in enumerate(zip(labels, fills)):
        label_box(s, xs[i], 1.95, widths[i] if i != 4 else 1.4, 1.1, [a, b], fill=f, size=15)
        if i:
            arrow(s, xs[i - 1] + (widths[i - 1] if i - 1 != 4 else 1.4), 2.5, xs[i], 2.5)
    rows = [["Placement", "What changes", "Matched control"],
            ["S1  KAN head", "h' = h + KAN(LN(h)) before dot-product scoring", "+MLP head"],
            ["S2  KAN input", "x' = x + KAN(LN(E[i])) before the GRU", "+MLP input"],
            ["S3  KAN-GRU cell", "gates stay linear + sigmoid; candidate n = tanh(KAN([x, r⊙h]))",
             "MLP-GRU cell"],
            ["Draft models", "first-draft ungated Pure KAN and Hybrid, re-run fairly", "Draft Hybrid (MLP)"]]
    table(s, 0.7, 3.45, 11.95, rows, col_w=[2.4, 6.85, 2.7], size=13, row_h=0.46, align_right_from=9)
    text(s, 0.7, 5.95, 11.95, 0.8, [{"runs": [
        ("KAN-NBR: ", {"bold": True, "color": KAN}),
        ("for baskets, an additive KAN over 7 repeat-purchase features (recency, count, share, decayed "
         "frequency, popularity, overdue ratio, in-history) gives one plottable curve per feature.", {})],
        "size": 15}])


def s_steps(prs, L):
    s = new_slide(prs, L, "4. Methodology: Steps and Deliverables")
    steps = [("Data pipelines", "YooChoose, Instacart, Dunnhumby with leakage checks", "dataset tables, EDA"),
             ("KAN module", "4 bases, residual block, matched MLP, unit tests", "kanrec/kan.py, 41 tests"),
             ("Baselines", "official GRU4Rec, basket-GRU, Pop, kNN, TopFreq, TIFU-KNN", "reproduction check"),
             ("KAN variants", "head, input, cell; MLP controls; KAN-NBR", "results per model and seed"),
             ("Evaluation", "Recall, MRR, NDCG, P@|B|, repeat/explore, Wilcoxon + Holm", "tables, figures"),
             ("Report", "notebook-generated tables in LaTeX, SRS, slides", "report, SRS, decks")]
    for i, (h, d, out) in enumerate(steps):
        r, c = divmod(i, 3)
        x, y = 0.7 + c * 4.05, 1.3 + r * 2.75
        box(s, x, y, 3.8, 2.5, fill=GREY_TINT)
        circle_num(s, x + 0.22, y + 0.22, i + 1, fill=BLUE)
        text(s, x + 0.8, y + 0.2, 2.85, 0.46, [{"text": h, "bold": True, "size": 17, "color": NAVY}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.25, y + 0.85, 3.35, 1.0, [{"text": d, "size": 14}])
        text(s, x + 0.25, y + 1.85, 3.35, 0.5, [{"runs": [("Deliverable: ", {"bold": True, "color": KAN}),
                                                           (out, {})], "size": 13}])


def s_timeline(prs, L):
    s = new_slide(prs, L, "4. Methodology: Timeline")
    tasks = [("Literature review", 1, 3, BLUE), ("Data acquisition and preprocessing", 2, 5, BLUE),
             ("KAN module and unit tests", 4, 6, BLUE), ("Baselines and reproduction", 5, 8, BLUE),
             ("KAN variants and MLP controls", 7, 11, KAN), ("Next-basket models and KAN-NBR", 9, 12, KAN),
             ("Ablations and GPU runs (3 seeds)", 11, 14, KAN), ("Analysis, report, SRS, slides", 12, 16, PURPLE)]
    x0, w_lab, weeks = 0.7, 3.6, 16
    gx, gw = x0 + w_lab, 12.65 - x0 - w_lab
    cw = gw / weeks
    for wk in range(weeks):
        text(s, gx + wk * cw, 1.25, cw, 0.35, [{"text": f"W{wk + 1}", "size": 11, "color": MUTED,
                                                 "align": PP_ALIGN.CENTER}])
    for i, (name, a, b, col) in enumerate(tasks):
        y = 1.7 + i * 0.6
        if i % 2 == 0:
            box(s, x0, y - 0.05, 12.65 - x0, 0.55, fill=GREY_TINT, shape=MSO_SHAPE.RECTANGLE)
        text(s, x0 + 0.1, y, w_lab - 0.2, 0.45, [{"text": name, "size": 14}], anchor=MSO_ANCHOR.MIDDLE)
        box(s, gx + (a - 1) * cw + 0.03, y + 0.07, (b - a + 1) * cw - 0.06, 0.31, fill=col, radius=0.3)
    for wk, lab, col in ((8, "Mid-term", ORANGE), (16, "End-term", PURPLE)):
        xx = gx + wk * cw - 0.02
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(xx), Inches(1.6), Inches(xx), Inches(6.55))
        ln.line.color.rgb = col
        ln.line.width = Pt(2)
        ln.line.dash_style = 4
        text(s, xx - 1.0, 6.55, 1.0 if wk == 16 else 2.0, 0.35,
             [{"text": lab, "bold": True, "size": 13, "color": col,
               "align": PP_ALIGN.RIGHT if wk == 16 else PP_ALIGN.CENTER}])


def s_working(prs, L, srs_name):
    s = new_slide(prs, L, "5. Working Model: Requirements and Architecture")
    text(s, 0.7, 1.2, 4.4, 0.45, [{"text": "Requirement analysis", "bold": True, "size": 17, "color": NAVY}])
    text(s, 0.7, 1.7, 4.4, 3.6, [
        {"text": f"Full SRS: {srs_name}", "bullet": True, "size": 14},
        {"text": "Functional: download, preprocess, train, evaluate and report every model from one "
                 "notebook", "bullet": True, "size": 14},
        {"text": "Hardware profiles: Mac (Apple MPS), UPES Blackwell GPU, Google Colab", "bullet": True,
         "size": 14},
        {"text": "Resumable: one result file per (dataset, model, seed)", "bullet": True, "size": 14},
        {"text": "Software: Python 3.12, PyTorch, pandas, Jupyter, LaTeX", "bullet": True, "size": 14},
    ], space_after=7)
    text(s, 5.5, 1.2, 7.2, 0.45, [{"text": "Technical diagram: pipeline", "bold": True, "size": 17,
                                    "color": NAVY}])
    nodes = [("Kaggle data", "kagglehub", 5.5, 1.8), ("Preprocessing", "data_session / data_basket", 9.15, 1.8),
             ("Models", "kan.py, session.py, basket.py, kannbr.py", 9.15, 3.35),
             ("Training + eval", "experiments.py, metrics.py", 5.5, 3.35),
             ("Results", "JSON + NPZ per run", 5.5, 4.9), ("Report", "report.py: LaTeX tables, figures", 9.15, 4.9)]
    for name, sub, x, y in nodes:
        label_box(s, x, y, 3.3, 1.05, [name, sub], fill=KANTINT if name == "Models" else TINT, size=15)
    arrow(s, 8.8, 2.33, 9.15, 2.33)
    arrow(s, 10.8, 2.85, 10.8, 3.35)
    arrow(s, 9.15, 3.88, 8.8, 3.88)
    arrow(s, 7.15, 4.4, 7.15, 4.9)
    arrow(s, 8.8, 5.43, 9.15, 5.43)
    box(s, 0.7, 5.45, 4.4, 1.3, fill=GREY_TINT)
    text(s, 0.9, 5.52, 4.0, 1.2, [{"text": "Code and results", "bold": True, "size": 14, "color": NAVY,
                                    "space": 3},
                                   {"text": REPO.replace("https://", ""), "size": 14, "color": MLP}])


def s_attained(prs, L):
    s = new_slide(prs, L, "5. Working Model: Attained Deliverables")
    yc = load("yoochoose64", "GRU4Rec__s42")["test"]["Recall@20"]
    stats = [("41", "unit tests passing"), ("3", "datasets with published protocols"),
             (f"{yc:.3f}", "GRU4Rec Recall@20 (SR-GNN table: 0.606)"),
             (str(len(list(RESULTS.glob('*/*.json')))), "finished model runs (resumable)")]
    for i, (big, lab) in enumerate(stats):
        x = 0.7 + i * 3.02
        box(s, x, 1.3, 2.8, 1.75, fill=TINT)
        text(s, x, 1.38, 2.8, 0.85, [{"text": big, "bold": True, "size": 36, "color": KAN,
                                      "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.15, 2.25, 2.5, 0.7, [{"text": lab, "size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}])
    text(s, 0.7, 3.35, 6, 0.45, [{"text": "Fixes to the first draft", "bold": True, "size": 17,
                                   "color": NAVY}])
    fixes = ["Leaky sample-level split replaced by session- and time-based splits",
             "Test set no longer used for model selection",
             "Degree-0 'B-spline' replaced by a verified cubic B-spline",
             "Ungated recurrent KAN replaced by gated KAN placements",
             "Every KAN paired with a parameter-matched MLP"]
    text(s, 0.7, 3.85, 6.0, 3.0, [{"text": t, "bullet": True, "size": 14} for t in fixes], space_after=6)
    pic = s.shapes.add_picture(str(ASSETS / "audit_bspline_basis.png"), Inches(7.0), Inches(3.5), width=Inches(5.65))
    text(s, 7.0, 3.6 + pic.height / 914400, 5.65, 0.4, [{"text": "First-draft basis (step functions) vs. corrected cubic B-spline",
                                     "size": 11, "color": MUTED, "align": PP_ALIGN.CENTER}])


def s_tests(prs, L):
    s = new_slide(prs, L, "6. Results: Test Cases and Reproduction")
    rows = [["Test case (pytest)", "Checks", "Status"],
            ["B-spline basis", "partition of unity; equals closed-form cubic", "Pass"],
            ["KAN layer", "shapes, gradients, parameter count formula", "Pass"],
            ["Matched MLP", "parameters within 2% of the KAN it controls", "Pass"],
            ["GRU4Rec wrapper", "identical outputs to the official model", "Pass"],
            ["Session iterator", "same batches as the official iterator", "Pass"],
            ["Metrics", "toy rankings; vectorised = loop version", "Pass"],
            ["Splits", "no session or user in two splits; time order", "Pass"],
            ["KAN-NBR", "batched scores = per-user scores", "Pass"]]
    table(s, 0.7, 1.25, 7.0, rows, col_w=[2.1, 4.0, 0.9], size=12, row_h=0.47, align_right_from=2,
          highlight={(r, 2): GOOD for r in range(1, len(rows))})
    text(s, 8.05, 1.2, 4.6, 0.45, [{"text": "Reproduction (YooChoose 1/64)", "bold": True, "size": 16,
                                     "color": NAVY}])
    rep = [("Item-kNN", "Item_kNN__s0", 0.5160), ("GRU4Rec", "GRU4Rec__s42", 0.6064)]
    cats, ours, pub = [], [], []
    for name, stem, p in rep:
        cats.append(name)
        ours.append(metric("yoochoose64", stem, "Recall@20"))
        pub.append(p)
    bar_chart(s, 8.0, 1.7, 4.7, 3.4, cats, [("Ours", ours), ("SR-GNN paper", pub)], colors=[MLP, BASE],
              legend=True, horizontal=False, vmin=0, vmax=0.8, font_size=12, number_format="0.000")
    text(s, 8.05, 5.2, 4.6, 1.5, [{"text": "Test Recall@20. Our GRU4Rec uses the tuned official settings "
                                            "(Hidasi & Karatzoglou 2018), so it exceeds the older number in "
                                            "the SR-GNN table.", "size": 12, "color": MUTED}])


def s_session_results(prs, L):
    s = new_slide(prs, L, "6. Results: Session-Based (YooChoose 1/64)")
    models = [("GRU4Rec", "GRU4Rec__s42", BASE), ("GRU4Rec (wide)", "GRU4Rec__wide__s42", BASE),
              ("+KAN head", "KAN_head__s42", KAN), ("+MLP head", "MLP_head__s42", MLP),
              ("+KAN input", "KAN_input__s42", KAN), ("+MLP input", "MLP_input__s42", MLP),
              ("KAN-GRU cell (RBF)", "KAN_GRU_cell__s42", KAN), ("MLP-GRU cell", "MLP_GRU_cell__s42", MLP),
              ("KAN-GRU cell (rational)", "KAN_GRU_cell__rational__G_8__s42", KAN),
              ("MLP-GRU cell (matched)", "MLP_GRU_cell__matched_to_rational__G_8__s42", MLP)]
    models = [m for m in models if load("yoochoose64", m[1])]
    bar_chart(s, 0.6, 1.2, 7.4, 5.6, [m[0] for m in models],
              [("Recall@20", [metric("yoochoose64", m[1], "Recall@20") for m in models])],
              point_colors=[m[2] for m in models], vmin=0.6, vmax=0.73, font_size=12, gap=45,
              title="Test Recall@20 (seed 42)")
    for i, (lab, col) in enumerate((("Baseline", BASE), ("KAN", KAN), ("Matched MLP", MLP))):
        box(s, 8.35 + i * 1.45, 1.35, 0.22, 0.22, fill=col, shape=MSO_SHAPE.RECTANGLE)
        text(s, 8.65 + i * 1.45, 1.29, 1.2, 0.35, [{"text": lab, "size": 12}])
    k = metric("yoochoose64", "KAN_GRU_cell__rational__G_8__s42", "Recall@20")
    m = metric("yoochoose64", "MLP_GRU_cell__matched_to_rational__G_8__s42", "Recall@20")
    g = metric("yoochoose64", "GRU4Rec__s42", "Recall@20")
    text(s, 8.35, 1.9, 4.35, 4.9, [
        {"text": "Findings", "bold": True, "size": 17, "color": NAVY, "space": 8},
        {"text": "KAN head and KAN input do not beat their MLP controls", "bullet": True, "size": 14},
        {"text": "Putting the non-linearity inside the GRU cell helps both KAN and MLP", "bullet": True,
         "size": 14},
        {"runs": [("Best: KAN-GRU cell with rational basis, ", {}), (f"{k:.4f}", {"bold": True, "color": KAN}),
                  (f" vs matched MLP {m:.4f} and GRU4Rec {g:.4f}", {})], "bullet": True, "size": 14},
        {"text": "A wider GRU4Rec with the same parameters does not help, so the gain is not just size",
         "bullet": True, "size": 14},
        {"text": "Seed 43 agrees; third seed pending", "bullet": True, "size": 14, "color": MUTED},
    ], space_after=8)


def s_basket_results(prs, L):
    s = new_slide(prs, L, "6. Results: Next-Basket (Instacart, Dunnhumby)")
    models = [("P-TopFreq", "P_TopFreq__s0", BASE), ("TIFU-KNN", "TIFU_KNN__s0", BASE),
              ("Basket-GRU", "Basket_GRU__s42", BASE), ("+KAN head", "KAN_head__s42", KAN),
              ("+MLP head", "MLP_head__s42", MLP), ("+KAN input", "KAN_input__s42", KAN),
              ("+MLP input", "MLP_input__s42", MLP), ("KAN-GRU cell", "KAN_GRU_cell__s42", KAN),
              ("MLP-GRU cell", "MLP_GRU_cell__s42", MLP), ("Logistic (features)", "Logistic__features__s42", BASE),
              ("KAN-NBR (additive)", "KAN_NBR__additive__s42", KAN)]
    for j, ds in enumerate(("instacart", "dunnhumby")):
        avail = [m for m in models if load(ds, m[1])]
        vals = [metric(ds, m[1], "Recall@10") for m in avail]
        bar_chart(s, 0.6 + j * 6.1, 1.2, 6.0, 4.75, [m[0] for m in avail], [("Recall@10", vals)],
                  point_colors=[m[2] for m in avail], vmin=0, vmax=0.42 if ds == "instacart" else 0.25,
                  font_size=11, gap=40, title=f"{ds.capitalize()}: test Recall@10")
    text(s, 0.7, 6.05, 11.95, 0.8, [{"runs": [
        ("Finding: ", {"bold": True, "color": KAN}),
        ("KAN head and input blocks hurt the basket-GRU; a KAN-GRU cell helps a little but its matched MLP "
         "helps as much or more. Frequency models win, and KAN-NBR matches the best of them while staying "
         "interpretable.", {})], "size": 14}])


def s_interpret(prs, L):
    s = new_slide(prs, L, "6. Results: Interpretability of KAN-NBR")
    s.shapes.add_picture(str(ASSETS / "kannbr_shapes.png"), Inches(0.6), Inches(1.15), width=Inches(12.1))
    text(s, 0.7, 5.45, 11.95, 1.4, [
        {"text": "Each curve is one learned KAN function: its contribution to the purchase logit.",
         "size": 14, "color": MUTED, "space": 6},
        {"runs": [("Overdue ratio: ", {"bold": True, "color": KAN}),
                  ("the purchase probability peaks when an item is about one usual interval overdue, then falls; a "
                   "logistic model can only fit a straight line here.", {})], "bullet": True, "size": 14},
        {"runs": [("Recency and frequency: ", {"bold": True, "color": KAN}),
                  ("recent and frequent items score higher, with diminishing returns.", {})], "bullet": True,
         "size": 14}], space_after=4)


def s_comparative(prs, L):
    s = new_slide(prs, L, "6. Results: KAN vs Matched MLP")
    pairs = [("YooChoose", "yoochoose64", "Head", "KAN_head__s42", "MLP_head__s42", "Recall@20"),
             ("YooChoose", "yoochoose64", "Input", "KAN_input__s42", "MLP_input__s42", "Recall@20"),
             ("YooChoose", "yoochoose64", "GRU cell (RBF)", "KAN_GRU_cell__s42", "MLP_GRU_cell__s42", "Recall@20"),
             ("YooChoose", "yoochoose64", "GRU cell (rational)", "KAN_GRU_cell__rational__G_8__s42",
              "MLP_GRU_cell__matched_to_rational__G_8__s42", "Recall@20"),
             ("Instacart", "instacart", "Head", "KAN_head__s42", "MLP_head__s42", "Recall@10"),
             ("Instacart", "instacart", "Input", "KAN_input__s42", "MLP_input__s42", "Recall@10"),
             ("Instacart", "instacart", "GRU cell", "KAN_GRU_cell__s42", "MLP_GRU_cell__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "Head", "KAN_head__s42", "MLP_head__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "Input", "KAN_input__s42", "MLP_input__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "GRU cell", "KAN_GRU_cell__s42", "MLP_GRU_cell__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "KAN-NBR vs MLP-NBR", "KAN_NBR__additive__s42", "MLP_NBR__matched__s42",
              "Recall@10")]
    rows = [["Dataset", "Placement", "Metric", "KAN", "MLP", "Winner"]]
    hl = {}
    for ds_name, ds, place, kst, mst, met in pairs:
        k, m = metric(ds, kst, met), metric(ds, mst, met)
        if k is None or m is None:
            rows.append([ds_name, place, met, "running" if k is None else f"{k:.4f}",
                         "running" if m is None else f"{m:.4f}", "-"])
            continue
        win = "KAN" if k > m + 0.002 else ("MLP" if m > k + 0.002 else "tie")
        rows.append([ds_name, place, met, f"{k:.4f}", f"{m:.4f}", win])
        hl[(len(rows) - 1, 5)] = KAN if win == "KAN" else (MLP if win == "MLP" else MUTED)
    table(s, 0.7, 1.2, 8.1, rows, col_w=[1.5, 2.4, 1.2, 1.0, 1.0, 1.0], size=12, row_h=0.43, highlight=hl,
          align_right_from=3)
    wins = sum(1 for r in rows[1:] if r[5] == "KAN")
    losses = sum(1 for r in rows[1:] if r[5] == "MLP")
    ties = sum(1 for r in rows[1:] if r[5] == "tie")
    for i, (big, lab, col) in enumerate(((wins, "KAN better", KAN), (losses, "MLP better", MLP),
                                         (ties, "within 0.002", MUTED))):
        y = 1.2 + i * 1.32
        box(s, 9.2, y, 3.45, 1.15, fill=GREY_TINT)
        text(s, 9.35, y, 1.0, 1.15, [{"text": str(big), "bold": True, "size": 34, "color": col}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, 10.4, y, 2.2, 1.15, [{"text": lab, "size": 15}], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.2, 5.25, 3.45, 1.5, [{"text": "Same parameter count in every pair; single seed (Mac profile), "
                                            "difference threshold 0.002.", "size": 12, "color": MUTED}])


def s_conclusion(prs, L):
    s = new_slide(prs, L, "7. Conclusion: Justification of Objectives")
    rows = [["Objective", "Outcome"],
            ["Fix first-draft errors", "Done: new splits, verified B-spline, gated placements, errata appendix"],
            ["KAN layer, four bases", "Done: unit-tested; rational basis is the most efficient"],
            ["KAN in GRU4Rec", "Head/input: no gain over MLP. Cell (rational): best Recall@20 on YooChoose"],
            ["KAN in basket-GRU", "No gain; frequency methods dominate next-basket accuracy"],
            ["Fair comparison", "Matched MLPs, wide GRU, significance tests, published protocols"],
            ["KAN-NBR", "Matches TIFU-KNN-level accuracy with readable feature curves"]]
    table(s, 0.7, 1.25, 11.95, rows, col_w=[3.2, 8.75], size=14, row_h=0.6, align_right_from=9)
    box(s, 0.7, 5.75, 11.95, 1.0, fill=KANTINT)
    text(s, 1.0, 5.78, 11.4, 0.95, [{"runs": [
        ("Take-away: ", {"bold": True, "color": KAN}),
        ("KAN layers are not a free accuracy gain. Where they help, an equally large MLP usually helps as much; "
         "their clearest value is interpretability on small, meaningful inputs.", {})], "size": 15}],
         anchor=MSO_ANCHOR.MIDDLE)


def s_future(prs, L):
    s = new_slide(prs, L, "7. Conclusion: Future Scope")
    items = [("GPU runs", "three seeds, yoochoose1/4 and full Instacart on the UPES Blackwell GPU"),
             ("Attention models", "group-rational KAN in the feed-forward block of SASRec"),
             ("Covariates", "KAN curves over time of day and days since last order"),
             ("More data", "Diginetica as a second session benchmark"),
             ("KAN training", "grid extension during training; faster fused kernels"),
             ("Deployment", "a small KAN-NBR service for interpretable basket suggestions")]
    for i, (h, d) in enumerate(items):
        r, c = divmod(i, 2)
        x, y = 0.7 + c * 6.05, 1.3 + r * 1.8
        box(s, x, y, 5.85, 1.6, fill=GREY_TINT)
        circle_num(s, x + 0.25, y + 0.25, i + 1, fill=PURPLE)
        text(s, x + 0.85, y + 0.22, 4.8, 0.46, [{"text": h, "bold": True, "size": 17, "color": NAVY}],
             anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.85, y + 0.75, 4.8, 0.8, [{"text": d, "size": 14}])


REFS = [
    "Hidasi, B., Karatzoglou, A., Baltrunas, L., & Tikk, D. (2016). Session-based recommendations with recurrent neural networks. ICLR.",
    "Hidasi, B., & Karatzoglou, A. (2018). Recurrent neural networks with top-k gains for session-based recommendations. CIKM, 843-852.",
    "Hidasi, B., & Czapp, A. T. (2023). The effect of third party implementations on reproducibility. RecSys, 272-282.",
    "Wu, S., Tang, Y., Zhu, Y., Wang, L., Xie, X., & Tan, T. (2019). Session-based recommendation with graph neural networks. AAAI, 346-353.",
    "Hu, H., He, X., Gao, J., & Zhang, Z.-L. (2020). Modeling personalized item frequency information for next-basket recommendation. SIGIR, 1071-1080.",
    "Li, M., Jullien, S., Ariannezhad, M., & de Rijke, M. (2023). A next basket recommendation reality check. ACM TOIS, 41(4), 116.",
    "van Maasakkers, L., Fok, D., & Donkers, B. (2023). Next-basket prediction in a high-dimensional setting using gated recurrent units. Expert Systems with Applications, 212, 118795.",
    "Liu, Z., Wang, Y., Vaidya, S., et al. (2025). KAN: Kolmogorov-Arnold networks. ICLR.",
    "Yu, R., Yu, W., & Wang, X. (2024). KAN or MLP: A fairer comparison. arXiv:2407.16674.",
    "Hou, Y., Ji, T., Zhang, D., & Stefanidis, A. (2025). Kolmogorov-Arnold networks: A critical assessment of claims, performance, and practical viability. arXiv:2407.11075.",
    "Genet, R., & Inzirillo, H. (2024). TKAN: Temporal Kolmogorov-Arnold networks. arXiv:2405.07344.",
    "Yang, X., & Wang, X. (2025). Kolmogorov-Arnold Transformer. ICLR.",
    "Park, J.-D., Kim, K.-M., & Shin, W.-Y. (2024). CF-KAN: Kolmogorov-Arnold network-based collaborative filtering. arXiv:2409.05878.",
    "Xu, Y., Fang, X., Gong, J., et al. (2026). Enhancing global and local interests fusion based on Kolmogorov-Arnold networks for sequential recommendation. Multimedia Systems, 32(2), 81.",
]


def s_references(prs, L, end_term):
    s = new_slide(prs, L, "References")
    paras = []
    if end_term:
        paras.append({"runs": [("A: Paper published. ", {"bold": True, "color": PURPLE}),
                               ("None yet; a paper on the KAN-GRU cell and KAN-NBR results is in preparation.", {})],
                      "size": 13, "space": 8})
    paras.append({"text": "B: Cited papers", "bold": True, "size": 13, "color": NAVY, "space": 4})
    paras += [{"text": r, "size": 11, "bullet": True} for r in REFS]
    text(s, 0.7, 1.15, 11.95, 5.8, paras, space_after=3)


def build(end_term):
    prs, L = open_template()
    title_slide, thanks = prs.slides[0], prs.slides[1]
    stage, col = ("End-Term Evaluation", PURPLE) if end_term else ("Mid-Term Evaluation", ORANGE)
    fill_title_slide(title_slide, stage, col)
    sections = ["Introduction", "Literature Review", "Objectives", "Methodology", "Working Model"]
    if end_term:
        sections += ["Results", "Conclusion"]
    sections += ["References"]
    s_contents(prs, L, sections)
    s_background(prs, L)
    s_concepts(prs, L)
    s_motivation(prs, L)
    s_application_data(prs, L)
    s_lit_table(prs, L)
    s_lit_inference(prs, L)
    s_swot(prs, L)
    s_objectives(prs, L)
    s_reference_model(prs, L)
    s_steps(prs, L)
    s_timeline(prs, L)
    s_working(prs, L, "KAN_Rec_SRS.docx")
    s_attained(prs, L)
    if end_term:
        s_tests(prs, L)
        s_session_results(prs, L)
        s_basket_results(prs, L)
        s_interpret(prs, L)
        s_comparative(prs, L)
        s_conclusion(prs, L)
        s_future(prs, L)
    s_references(prs, L, end_term)
    move_to_end(prs, thanks)
    out = PRES / f"KAN_Rec_{'EndTerm' if end_term else 'MidTerm'}.pptx"
    prs.save(out)
    print("wrote", out, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    build(False)
    build(True)
