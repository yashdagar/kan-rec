"""Builds the Software Requirements Specification on the UPES SRS template.

    uv run --with python-docx --with pypdf python tools/build_srs.py

Diagrams are rendered from presentation/diagrams/*.puml with PlantUML. Page numbers in the table of
contents are filled from a LibreOffice render of a first pass.
"""
import datetime as dt
import json
import re
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
PRES = ROOT / "presentation"
TEMPLATE = PRES / "templates" / "SRS Template.docx"
DIAG = PRES / "diagrams"
ASSETS = PRES / "assets"
RESULTS = ROOT / "results" / "mac"
OUT = PRES / "KAN_Rec_SRS.docx"
REPO = "https://github.com/yashdagar/kan-rec"
TITLE = "Kolmogorov-Arnold Networks in Recurrent Recommenders: Session-Based and Next-Basket Prediction"
DATE = dt.date(2026, 10, 6)
STUDENTS = [("B.Tech CSE", "500124297", "Abhishek Yadav"),
            ("B.Tech CSE", "500121969", "Saksham Agrawal"),
            ("B.Tech CSE", "500125147", "Yash Dagar")]
TNR = "Times New Roman"


def res(ds, stem):
    p = RESULTS / ds / f"{stem}.json"
    return json.loads(p.read_text()) if p.exists() else None


def _run_font(run, size=12, bold=False, italic=False, underline=False, color=None):
    run.font.name = TNR
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), TNR)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    if color:
        run.font.color.rgb = color


def _para_format(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, line=1.15, keep_next=False):
    pf = p.paragraph_format
    pf.alignment = align
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line
    pf.keep_with_next = keep_next


class Writer:
    def __init__(self, doc, anchor):
        self.doc = doc
        self.anchor = anchor
        self.fig = 0
        self.tab = 0
        self.headings = {}

    def _new_p(self):
        p = OxmlElement("w:p")
        self.anchor.addprevious(p)
        para = Paragraph(p, self.doc._body)
        para.style = self.doc.styles["Normal"]
        return para

    def para(self, text="", size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=6,
             runs=None, keep_next=False, before=0):
        p = self._new_p()
        _para_format(p, align=align, after=after, keep_next=keep_next, before=before)
        for t, o in (runs or [(text, {})]):
            r = p.add_run(t)
            _run_font(r, o.get("size", size), o.get("bold", bold), o.get("italic", italic), o.get("underline", False))
        return p

    def heading(self, text, key=None):
        p = self.para(text.upper(), size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=8, before=12,
                      keep_next=True)
        if key:
            self.headings[key] = text
        return p

    def sub(self, text, key=None):
        p = self.para(align=WD_ALIGN_PARAGRAPH.LEFT, after=4, before=8, keep_next=True,
                      runs=[(text, {"bold": True, "underline": True})])
        if key:
            self.headings[key] = text
        return p

    def minor(self, text):
        return self.para(align=WD_ALIGN_PARAGRAPH.LEFT, after=3, before=4, keep_next=True,
                         runs=[(text, {"bold": True, "italic": True})])

    def bullets(self, items, numbered=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        for i, it in enumerate(items):
            p = self._new_p()
            _para_format(p, after=3, align=align)
            pf = p.paragraph_format
            pf.left_indent = Inches(0.35)
            pf.first_line_indent = Inches(-0.22)
            mark = f"{i + 1}." if numbered else "•"
            r = p.add_run(f"{mark}\t")
            _run_font(r)
            tabs = pf.tab_stops
            tabs.add_tab_stop(Inches(0.35))
            if isinstance(it, tuple):
                r = p.add_run(it[0])
                _run_font(r, bold=True)
                r = p.add_run(it[1])
                _run_font(r)
            else:
                r = p.add_run(it)
                _run_font(r)

    def page_break(self):
        p = self._new_p()
        p.add_run().add_break(WD_BREAK.PAGE)

    def caption(self, text, kind):
        if kind == "Table":
            self.tab += 1
            n = self.tab
        else:
            self.fig += 1
            n = self.fig
        return self.para(align=WD_ALIGN_PARAGRAPH.CENTER, after=8, keep_next=(kind == "Table"),
                         runs=[(f"{kind} {n}: ", {"bold": True, "size": 11}), (text, {"size": 11})])

    def figure(self, path, caption, width=6.0, source="Authors' own work"):
        p = self._new_p()
        _para_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, keep_next=True)
        p.add_run().add_picture(str(path), width=Inches(width))
        self.caption(f"{caption} (Source: {source})", "Figure")

    def table(self, rows, widths, caption=None, header=True, align_right=()):
        if caption:
            self.caption(caption, "Table")
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        self.anchor.addprevious(t._tbl)
        t.style = self.doc.styles["Table Grid"] if "Table Grid" in [s.name for s in self.doc.styles] else None
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Inches(widths[j])
                p = cell.paragraphs[0]
                _para_format(p, align=WD_ALIGN_PARAGRAPH.RIGHT if (j in align_right and i > 0) else
                             WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
                r = p.add_run(str(val))
                _run_font(r, 10, bold=(header and i == 0))
                if header and i == 0:
                    _shade(cell, "D9E2F3")
        for j, w in enumerate(widths):
            t.columns[j].width = Inches(w)
        spacer = self._new_p()
        _para_format(spacer, after=4)
        return t


def _shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def _borders(table):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)


def _set_cell_text(cell, text, size=12, bold=False, align=None):
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run(text)
    _run_font(r, size, bold)
    if align is not None:
        p.alignment = align


def _replace_para_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def fill_front_matter(doc, page_numbers):
    for p in doc.paragraphs:
        if p.text.strip() == "<Project>":
            _replace_para_text(p, TITLE)
        elif p.text.strip() == "<Date>":
            _replace_para_text(p, DATE.strftime("%d %B %Y").lstrip("0"))
    cover, toc, rev = doc.tables[0], doc.tables[1], doc.tables[2]
    for i, (spec, sap, name) in enumerate(STUDENTS):
        for j, v in enumerate((spec, sap, name)):
            _set_cell_text(cover.rows[i + 1].cells[j], v, 12, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in list(cover.rows)[len(STUDENTS) + 1:]:
        row._tr.getparent().remove(row._tr)
    for row in list(toc.rows)[1:]:
        cells = row.cells
        parts = []
        for c in cells[:-1]:
            t = c.text.strip()
            if t and t not in parts:
                parts.append(t)
        num = page_numbers.get(_norm(" ".join(parts)))
        if num:
            _set_cell_text(cells[-1], str(num), 10, align=WD_ALIGN_PARAGRAPH.CENTER)
    rv = [(DATE.strftime("%d-%m-%Y"), "Version 1.0: complete SRS", "End-term submission of the major project", "")]
    for i, vals in enumerate(rv):
        for j, v in enumerate(vals):
            _set_cell_text(rev.rows[i + 1].cells[j], v, 10)


def _norm(s):
    return re.sub(r"[^a-z]", "", s.lower().replace("data structure", "").replace("characteristic of data", ""))


def strip_instructions(doc):
    body = doc.element.body
    children = list(body.iterchildren())
    rev = doc.tables[2]._tbl
    guide = doc.tables[3]._tbl
    start = children.index(rev) + 1
    end = children.index(guide)
    for el in children[start:end + 1]:
        body.remove(el)
    anchor = body.find(qn("w:sectPr"))
    p = OxmlElement("w:p")
    anchor.addprevious(p)
    Paragraph(p, doc._body).add_run().add_break(WD_BREAK.PAGE)
    return anchor


def write_content(w):
    def m(ds, stem, k="Recall@20"):
        r = res(ds, stem)
        return f"{r['test'][k]:.4f}" if r else "n/a"

    w.heading("1 Introduction")
    w.sub("1.1 Purpose of the Project", "Purpose of the Project")
    w.para("Recommender systems rank the items of a large catalogue for a user. This project studies two "
           "sequential tasks: session-based recommendation, which predicts the next item clicked in an anonymous "
           "browsing session, and next-basket recommendation, which predicts the products in a customer's next "
           "grocery basket from their earlier baskets. Both tasks are usually modelled with gated recurrent units "
           "(GRUs): GRU4Rec for sessions and the basket-GRU of van Maasakkers et al. (2023) for baskets.")
    w.para("Kolmogorov-Arnold Networks (KANs) replace the fixed activation functions of a multi-layer perceptron "
           "(MLP) with a learnable univariate function on every edge. They are claimed to be more accurate per "
           "parameter and more interpretable than MLPs, but later studies found that these advantages mostly "
           "disappear when the two are compared at the same size.")
    w.para(runs=[("Problem statement. ", {"bold": True}),
                 ("Determine whether, and where, KAN layers improve GRU-based session and next-basket recommenders, "
                  "under a protocol that allows a fair conclusion: standard leakage-free splits, verified "
                  "implementations, parameter-matched MLP controls, and evaluation on three public datasets.", {})])
    w.para(runs=[("Motivation. ", {"bold": True}),
                 ("The first draft of this project compared KAN and GRU models on a data split that leaked test "
                  "sessions into training, selected models on the test set and used a B-spline routine that "
                  "produced step functions. Its numbers could not support any conclusion. This system re-does the "
                  "study so that both a positive and a negative result are trustworthy.", {})])

    w.sub("1.2 Target Beneficiary", "Target Beneficiary")
    w.bullets([("Recommender-system researchers: ", "a controlled, reproducible answer to whether KAN layers help "
                                                      "recurrent recommenders, with all code and results public."),
               ("E-commerce and online-grocery platforms: ", "evidence on which model to deploy, and an interpretable "
                                                              "repeat-purchase model (KAN-NBR) whose behaviour can be "
                                                              "explained to business users."),
               ("Students and educators: ", "a worked example of a fair model comparison, including the errors of a "
                                             "first draft and how they were fixed."),
               ("End users (shoppers): ", "indirectly, through better and more explainable recommendations that need "
                                          "no long-term tracking in the session setting.")])

    w.sub("1.3 Project Scope", "Project Scope")
    w.para("The software is a research pipeline, not an end-user application. It downloads the datasets, "
           "preprocesses them with published protocols, trains every model variant, evaluates it once on a held-out "
           "test set, runs significance tests, and writes the tables and figures of the report. Its goals are:")
    w.bullets(["a correct, unit-tested KAN layer with four basis families (B-spline, Gaussian RBF, Chebyshev, "
               "group-rational);",
               "KAN layers in three placements (head, input, GRU cell) of the official GRU4Rec and of the basket-GRU, "
               "each with a parameter-matched MLP control;",
               "strong non-neural baselines (Pop, S-Pop, item-kNN, TopFreq variants, TIFU-KNN);",
               "KAN-NBR, an additive KAN over interpretable repeat-purchase features;",
               "three hardware profiles (Mac, UPES Blackwell GPU, Google Colab) selected by one variable."])
    w.para("Deliverables: the notebook kan_rec.ipynb and the UPES notebook kan_rec_upes.ipynb; the Python package "
           "kanrec with 41 unit tests; per-run result files; the project report (LaTeX and a single-file Overleaf "
           f"version); the mid-term and end-term presentations; and this SRS. All are public at {REPO}. "
           "Out of scope: a production service, a graphical user interface and real-time serving.")

    w.sub("1.4 References", "References")
    refs = [
        "Hidasi, B., Karatzoglou, A., Baltrunas, L., & Tikk, D. (2016). Session-based recommendations with recurrent neural networks. ICLR.",
        "Hidasi, B., & Karatzoglou, A. (2018). Recurrent neural networks with top-k gains for session-based recommendations. CIKM, 843-852.",
        "Hidasi, B., & Czapp, A. T. (2023). The effect of third party implementations on reproducibility. RecSys, 272-282.",
        "Wu, S., Tang, Y., Zhu, Y., Wang, L., Xie, X., & Tan, T. (2019). Session-based recommendation with graph neural networks. AAAI, 346-353.",
        "Hu, H., He, X., Gao, J., & Zhang, Z.-L. (2020). Modeling personalized item frequency information for next-basket recommendation. SIGIR, 1071-1080.",
        "Li, M., Jullien, S., Ariannezhad, M., & de Rijke, M. (2023). A next basket recommendation reality check. ACM TOIS, 41(4), 116.",
        "van Maasakkers, L., Fok, D., & Donkers, B. (2023). Next-basket prediction in a high-dimensional setting using gated recurrent units. Expert Systems with Applications, 212, 118795.",
        "Liu, Z., Wang, Y., Vaidya, S., et al. (2025). KAN: Kolmogorov-Arnold networks. ICLR.",
        "Yu, R., Yu, W., & Wang, X. (2024). KAN or MLP: A fairer comparison. arXiv:2407.16674.",
        "Genet, R., & Inzirillo, H. (2024). TKAN: Temporal Kolmogorov-Arnold networks. arXiv:2405.07344.",
        "Yang, X., & Wang, X. (2025). Kolmogorov-Arnold Transformer. ICLR.",
        "IEEE Std 830-1998, IEEE Recommended Practice for Software Requirements Specifications.",
        "Official GRU4Rec (PyTorch): https://github.com/hidasib/GRU4Rec_PyTorch_Official",
        "YooChoose (RecSys Challenge 2015): https://www.kaggle.com/datasets/chadgostopp/recsys-challenge-2015",
        "Instacart Market Basket Analysis: https://www.kaggle.com/c/instacart-market-basket-analysis",
        "Dunnhumby, The Complete Journey: https://www.kaggle.com/datasets/frtgnn/dunnhumby-the-complete-journey",
        "Next-basket GRU code: https://github.com/luukvanmaasakkers/nextbasketpredictionGRU",
        "TARS / tbp-next-basket: https://github.com/GiulioRossetti/tbp-next-basket",
        f"Project repository: {REPO}",
    ]
    w.bullets(refs, numbered=True, align=WD_ALIGN_PARAGRAPH.LEFT)

    w.page_break()
    w.heading("2 Project Description")
    w.sub("2.1 Reference Algorithm", "Reference Algorithm")
    w.para("The reference algorithm is GRU4Rec (Hidasi et al. 2016, 2018) in its official PyTorch implementation. "
           "KAN layers are added to it, and to the basket-GRU, without changing the training procedure.")
    w.minor("Algorithm 1: GRU4Rec training with KAN blocks")
    w.bullets(["Sort training clicks by session and time; keep an offset array marking where each session starts.",
               "Fill 48 mini-batch slots with the first 48 sessions; each slot advances one click per step.",
               "For the current items, look up embeddings x = E[i]; optionally apply the input block x' = x + f(LN(x)).",
               "Update the hidden state with the GRU cell (or the KAN-GRU cell, whose candidate state is tanh(KAN([x, r * h]))).",
               "Optionally apply the head block h' = h + f(LN(h)), then score the targets and 2048 shared negative items "
               "sampled by popularity^0.2 with the dot product E h'.",
               "Minimise softmax cross-entropy with logQ correction using Adagrad (lr 0.07; added parameters lr 0.001).",
               "When a session ends, reset its slot's hidden state and load the next session.",
               "After each epoch compute validation Recall@20 and keep the best weights; evaluate once on the test set."],
              numbered=True)
    w.para("The KAN layer computes, for each output o, y_o = sum over inputs i of phi_oi(x_i), where "
           "phi(x) = w * silu(x) + sum_k c_k B_k(x) and B_k is a fixed basis (cubic B-spline by the Cox-de Boor "
           "recursion, Gaussian RBF, Chebyshev polynomials of tanh(x), or a group-rational function). Each KAN block "
           "has an MLP control with the same number of parameters (within 2%).")
    w.minor("Data structures")
    w.table([["Structure", "Used for", "Implementation"],
             ["Offset array + item index array", "session-parallel mini-batches", "NumPy int64 arrays, on GPU blocks"],
             ["Embedding matrix E (items x 480)", "shared input and output item vectors", "torch.nn.Parameter"],
             ["List of baskets per user", "basket histories", "Python lists of int arrays (BasketData)"],
             ["Index + offset lists", "multi-hot baskets without dense vectors", "torch.nn.EmbeddingBag"],
             ["Sparse co-occurrence matrix", "item-kNN and product clustering", "SciPy CSR matrix"],
             ["Feature matrix (pairs x 7)", "KAN-NBR training pairs", "float32 tensor"],
             ["Result record", "metrics, history, timing per run", "JSON file + NPZ per-user arrays"]],
            [2.2, 2.2, 2.0], "Data structures used by the system.")

    w.sub("2.2 Characteristic of Data", "Characteristic of Data")
    w.para("All data are secondary, public datasets downloaded from Kaggle; no primary data are collected. "
           f"Table {w.tab + 1} summarises them after preprocessing.")
    w.table([["Property", "YooChoose 1/64", "Instacart", "Dunnhumby"],
             ["Source", "RecSys Challenge 2015", "Instacart Market Basket", "The Complete Journey"],
             ["Raw size", "33.0 M clicks, 9.25 M sessions", "3.4 M orders, 206,209 users", "2.60 M rows, 2,500 households"],
             ["Unit", "click in a session", "basket of a user", "basket of a household"],
             ["After preprocessing", "139,628 sessions", "20,620 users (10% sample)", "2,478 households"],
             ["Items", "17,264", "12,058 product groups", "10,370"],
             ["Mean length", "4.1 clicks", "16.2 baskets, 10.1 items", "66.2 baskets, 8.7 items"],
             ["Test targets", "55,357 next clicks", "10,310 baskets", "1,239 baskets"],
             ["Repeat share of target", "-", "0.60", "0.53"]],
            [1.5, 1.7, 1.7, 1.6], "Datasets after preprocessing (Mac profile).")
    w.minor("Sampling")
    w.bullets(["YooChoose: the most recent 1/64 of training sessions (standard yoochoose1/64); sessions of the last day "
               "form the test set and the last day of training data the validation set.",
               "Instacart: a random 10% of users on the Mac profile (all users on the GPU profile); the last basket of "
               "each user is the target; users are split 50/50 into validation and test with a fixed seed.",
               "Dunnhumby: all households with at least 3 baskets, the 100 most recent baskets kept; same 50/50 split."])
    w.minor("Statistical processing")
    w.bullets(["Support filters: items in at least 5 sessions (YooChoose) or bought at least 50 times (Dunnhumby).",
               "Instacart rare products (fewer than 500 purchases) merged greedily within their aisle by cosine "
               "similarity of purchase vectors.",
               "Paired Wilcoxon signed-rank tests on per-prediction or per-user metrics with Holm correction."])

    w.sub("2.3 SWOT Analysis", "SWOT Analysis")
    w.table([["Strengths", "Weaknesses"],
             ["Official GRU4Rec code and published protocols; parameter-matched controls; 41 unit tests; three "
              "datasets and two tasks; resumable runs.",
              "KAN layers are slower to train; extra hyperparameters (basis, grid size); Mac results use one seed and "
              "a 10% Instacart sample."],
             ["Opportunities", "Threats"],
             ["Interpretable repeat-purchase models (KAN-NBR); larger runs on the UPES Blackwell GPU; KAN layers in "
              "attention-based recommenders.",
              "KANs may show no advantage over MLPs; transformer recommenders dominate recent work; limited compute "
              "time for many variants."]],
            [3.2, 3.2], "SWOT analysis.")
    w.para("Justification: the strengths address the exact weaknesses of the first draft (leakage, unverified code, "
           "no capacity control), so whatever the outcome, the comparison is credible. The main threat, a negative "
           "result for KAN, is mitigated by framing a rigorous negative result as a contribution and by KAN-NBR, "
           "which uses KANs where their interpretability matters.")

    w.sub("2.4 Project Features", "Project Features")
    w.bullets([("Profile switch: ", "one variable selects Mac, UPES or Colab settings, paths and device."),
               ("Data pipelines: ", "download, preprocessing and leakage assertions for three datasets."),
               ("KAN module: ", "four bases, residual block with zero-initialised output, matched MLP."),
               ("Model zoo: ", "GRU4Rec and basket-GRU with KAN or MLP head, input or cell; draft models; KAN-NBR."),
               ("Baselines: ", "Pop, S-Pop, item-kNN, G/P/GP-TopFreq, last basket, TIFU-KNN."),
               ("Evaluation: ", "Recall, MRR, NDCG, PHR, P@|B|, average rank, repeat/explore split."),
               ("Resumable experiments: ", "one result file per dataset, model and seed."),
               ("Reporting: ", "LaTeX tables, figures, significance tests, learned-function plots.")])
    w.figure(DIAG / "usecase.png", "Level-2 use case diagram", 6.0)

    w.sub("2.5 User Classes and Characteristics", "User Classes and Characteristics")
    w.table([["User class", "Characteristics", "Main use"],
             ["Researcher (project team)", "Python and PyTorch skills; runs the notebook on all profiles",
              "run experiments, extend models"],
             ["Mentor / evaluator", "reads results; may re-run selected cells", "review report and results"],
             ["External researcher", "clones the public repository", "reproduce or extend the study"],
             ["Practitioner", "data scientist at a retailer", "reuse KAN-NBR or baselines on own data"]],
            [1.8, 2.7, 1.9], "User classes.")

    w.sub("2.6 Design and Implementation Constraints", "Design and Implementation Constraints")
    w.bullets([("Hardware: ", "Apple M1 Pro with 16 GB unified memory (Mac profile); PyTorch MPS has no float64, so "
                               "all tensors are float32; memory is freed after every run. UPES profile: NVIDIA "
                               "Blackwell GPU with CUDA. Colab: T4/L4 GPU with session time limits."),
               ("Timing: ", "one GRU4Rec epoch on YooChoose 1/64 takes about 75 s on the Mac; the full Mac run "
                            "takes about two days, so every run must be resumable."),
               ("Interfaces: ", "official GRU4Rec code (academic licence, cloned at a pinned commit, not "
                                "redistributed); kagglehub for downloads."),
               ("Technologies: ", "Python 3.12, PyTorch 2.x, NumPy, pandas, SciPy, Matplotlib, Jupyter, uv, LaTeX."),
               ("Parallel operation: ", "session-parallel mini-batches; GPU kernels; chunked evaluation to bound "
                                        "memory."),
               ("Standards: ", "PEP 8 style; unit tests with pytest; fixed random seeds; no test-set use for model "
                               "selection."),
               ("Security: ", "the Kaggle API token stays in the user's home directory and is never committed.")])

    w.sub("2.7 Design Diagrams", "Design diagrams")
    w.figure(DIAG / "class.png", "Class diagram of the kanrec package", 6.2)
    w.figure(DIAG / "activity.png", "Activity diagram of one notebook run", 3.4)
    w.figure(DIAG / "sequence.png", "Sequence diagram for training one session model", 6.3)
    w.figure(DIAG / "dfd.png", "Level-1 data flow diagram", 2.6)
    w.figure(DIAG / "state.png", "State diagram of one experiment run", 4.4)
    w.figure(DIAG / "collaboration.png", "Collaboration (communication) diagram for a KAN-head run", 4.4)
    w.figure(DIAG / "deployment.png", "Deployment diagram", 6.3)

    w.sub("2.8 Assumption and Dependencies", "Assumption and Dependencies")
    w.bullets(["The Kaggle datasets remain available with the same content; Instacart requires accepting the "
               "competition rules once.",
               "The official GRU4Rec repository remains available at the pinned commit.",
               "Published numbers (SR-GNN table, van Maasakkers et al.) are correct and use the protocols described.",
               "Test sets are large enough that single-seed differences above about 0.002 are meaningful; this is "
               "checked with significance tests and extra seeds.",
               "UPES HPC access and GPU time are available for the multi-seed runs."])

    w.heading("3 System Requirements")
    w.sub("3.1 User Interface", "User Interface")
    w.para("The system has no graphical application. Its user interface is the Jupyter notebook:")
    w.bullets(["a configuration cell with PROFILE = \"mac\" | \"upes\" | \"colab\";",
               "progress logs per epoch (loss, validation metrics, seconds per epoch), also written to results/<profile>/run.log;",
               "inline tables and figures (results, significance, curves, learned KAN functions);",
               "command-line entry points: uv run pytest, tools/build_notebook.py, tools/flatten_report.py;",
               "the compiled PDF report and the presentations for evaluators."])
    w.sub("3.2 Software Interface", "Software Interface")
    w.table([["Module", "Provides", "Used by"],
             ["kan.py", "KANLinear, ResidualBlock, matched_mlp", "session.py, basket.py, kannbr.py"],
             ["data_session.py", "load_yoochoose_clicks, prepare_yoochoose", "notebook"],
             ["data_basket.py", "load_instacart, load_dunnhumby, BasketData", "notebook, basket.py, kannbr.py"],
             ["session.py", "KANGRU4RecModel, KAN-GRU cell, train_session_model, evaluate", "experiments.py"],
             ["basket.py", "BasketGRU, train_basket_gru, TopFreq, TIFU-KNN", "experiments.py"],
             ["kannbr.py", "NBRFeaturizer, KANNBR, train_kannbr, shape_functions", "experiments.py, notebook"],
             ["metrics.py", "session and basket metrics", "session.py, basket.py, kannbr.py"],
             ["experiments.py", "run_* functions, load_results", "notebook"],
             ["report.py", "LaTeX tables, Wilcoxon + Holm", "notebook"]],
            [1.5, 3.1, 1.9], "Module interfaces.")
    w.para("Modules communicate through in-process Python calls; data are passed as NumPy arrays, pandas frames and "
           "PyTorch tensors. Each run_* function takes the dataset, a model name and a seed, returns a result "
           "dictionary, and writes it to results/<profile>/<dataset>/<model>__s<seed>.json. External services are "
           "reached only through kagglehub (downloads) and git (code and results).", align=WD_ALIGN_PARAGRAPH.LEFT)
    w.sub("3.3 Database Interface", "Database Interface")
    w.para("No database management system is used. Persistent data are files: raw CSV files in data/ (git-ignored), "
           "cached preprocessed splits, model weights (.pt, git-ignored), and one JSON + NPZ pair per run. The JSON "
           "record holds the dataset, model, seed, configuration, per-epoch history, best epoch, validation and "
           "test metrics, parameter count and timing; the NPZ holds per-prediction or per-user metric values for "
           "significance tests. File existence is the resume key.")
    w.sub("3.4 Protocols", "Protocols")
    w.bullets(["HTTPS with an API token for Kaggle downloads (kagglehub); about 2 GB in total.",
               "Git over HTTPS for GitHub; results are pushed after each batch of runs.",
               "No network access during training or evaluation; runs are synchronised only through result files.",
               "On Colab, results are written to Google Drive so an interrupted session resumes."])

    w.heading("4 Non-functional Requirements")
    w.sub("4.1 Performance Requirements", "Performance requirements")
    rows = [["Model (Mac, seed 42)", "Parameters", "s / epoch", "Test metric"]]
    for ds, stem, label, k in [("yoochoose64", "GRU4Rec__s42", "GRU4Rec", "Recall@20"),
                               ("yoochoose64", "KAN_GRU_cell__rational__G_8__s42", "KAN-GRU cell (rational)", "Recall@20"),
                               ("yoochoose64", "MLP_GRU_cell__matched_to_rational__G_8__s42", "MLP-GRU cell (matched)", "Recall@20"),
                               ("yoochoose64", "KAN_GRU_cell__bspline__G_8__s42", "KAN-GRU cell (B-spline)", "Recall@20"),
                               ("dunnhumby", "Basket_GRU__s42", "Basket-GRU (Dunnhumby)", "Recall@10"),
                               ("dunnhumby", "KAN_NBR__additive__s42", "KAN-NBR (Dunnhumby)", "Recall@10"),
                               ("instacart", "Basket_GRU__s42", "Basket-GRU (Instacart 10%)", "Recall@10")]:
        r = res(ds, stem)
        if r:
            rows.append([label, f"{r['n_params']:,}", f"{r['train_s_per_epoch']:.1f}", f"{k} {r['test'][k]:.4f}"])
    w.table(rows, [2.4, 1.3, 1.0, 1.8], "Measured cost and accuracy of key models.", align_right=(1, 2))
    w.bullets(["A YooChoose 1/64 epoch must finish within 2 minutes on the Mac and evaluation of all 55,357 test "
               "predictions within 5 seconds.",
               "Peak memory must stay below 12 GB on the 16 GB Mac (chunked evaluation, memory freed after runs).",
               "Every KAN model must report parameters, seconds per epoch and inference time next to accuracy.",
               "The UPES profile must finish three seeds of the key models within one day of GPU time."])
    w.sub("4.2 Security Requirements", "Security requirements")
    w.bullets(["The datasets are anonymised by their publishers; no personal data are collected or linked.",
               "Raw data are not redistributed: data/ is git-ignored, in line with the Kaggle terms.",
               "Credentials (Kaggle token, GitHub token) are never stored in the repository or the notebooks.",
               "The official GRU4Rec code is cloned, not committed, respecting its licence.",
               "Validation: unit tests, leakage assertions in every pipeline, and a reproduction check against "
               "published results."])
    w.sub("4.3 Software Quality Attributes", "Software Quality Attributes")
    w.table([["Attribute", "How it is met"],
             ["Correctness", "41 unit tests; wrapper output identical to official GRU4Rec; split assertions"],
             ["Reliability / robustness", "resumable runs; zero-initialised residual blocks; stable learning rates"],
             ["Availability", "public GitHub repository; Kaggle data; Colab profile needs no local hardware"],
             ["Adaptability / flexibility", "new bases, placements or datasets added as one function or dictionary entry"],
             ["Interoperability", "standard file formats (JSON, NPZ, CSV, LaTeX, PDF)"],
             ["Maintainability", "small modules with one role each; notebooks generated from one builder script"],
             ["Portability", "runs on macOS (MPS), Linux (CUDA) and Colab; pure Python"],
             ["Reusability", "KAN layer and metrics usable independently of the notebook"],
             ["Testability", "pure functions for metrics and features; pytest suite"],
             ["Usability", "one profile variable; one command to run; tables written directly into the report"]],
            [2.0, 4.4], "Software quality attributes.")

    w.heading("5 Other Requirements")
    w.bullets([("Reproducibility: ", "fixed seeds, pinned third-party commit, recorded configuration in every result."),
               ("Fairness of comparison: ", "every KAN variant has a parameter-matched MLP control; the decision rule "
                                            "is applied on validation data only."),
               ("Documentation: ", "README with instructions per profile; report with errata appendix."),
               ("Ethics: ", "results are reported whether or not they favour KANs.")])

    w.page_break()
    w.heading("Appendix A: Glossary")
    w.table([["Term", "Meaning"],
             ["KAN", "Kolmogorov-Arnold Network: learnable univariate function on every edge"],
             ["MLP", "multi-layer perceptron with fixed activations"],
             ["GRU / GRU4Rec", "gated recurrent unit / GRU-based session recommender"],
             ["Basket-GRU", "GRU over multi-hot baskets with a sigmoid output per product"],
             ["KAN-NBR", "additive KAN over repeat-purchase features for next-basket recommendation"],
             ["RBF, B-spline, Chebyshev, GR-KAN", "basis families used inside KAN layers"],
             ["Recall@K / HR@K", "share of targets ranked in the top K"],
             ["MRR@K", "mean reciprocal rank, zero when the target is outside the top K"],
             ["NDCG@K", "normalised discounted cumulative gain"],
             ["P@|B|", "precision at the size of the true basket"],
             ["TIFU-KNN", "time-decayed item frequency with nearest-neighbour users"],
             ["MPS", "Apple Metal Performance Shaders backend of PyTorch"],
             ["Wilcoxon, Holm", "paired signed-rank test; multiple-comparison correction"],
             ["SRS", "Software Requirements Specification"]],
            [2.0, 4.4])
    w.heading("Appendix B: Analysis Model")
    w.para("The analysis model is a controlled comparison. For each dataset, every variant shares the baseline's "
           "hyperparameters; a KAN placement is called helpful only if, on validation data, it beats both the "
           f"baseline and its parameter-matched MLP. Table {w.tab + 1} lists the current test results for these pairs.")
    pairs = [("YooChoose", "yoochoose64", "Head", "KAN_head__s42", "MLP_head__s42", "Recall@20"),
             ("YooChoose", "yoochoose64", "Input", "KAN_input__s42", "MLP_input__s42", "Recall@20"),
             ("YooChoose", "yoochoose64", "Cell (RBF)", "KAN_GRU_cell__s42", "MLP_GRU_cell__s42", "Recall@20"),
             ("YooChoose", "yoochoose64", "Cell (rational)", "KAN_GRU_cell__rational__G_8__s42",
              "MLP_GRU_cell__matched_to_rational__G_8__s42", "Recall@20"),
             ("Instacart", "instacart", "Head", "KAN_head__s42", "MLP_head__s42", "Recall@10"),
             ("Instacart", "instacart", "Input", "KAN_input__s42", "MLP_input__s42", "Recall@10"),
             ("Instacart", "instacart", "Cell", "KAN_GRU_cell__s42", "MLP_GRU_cell__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "Head", "KAN_head__s42", "MLP_head__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "Input", "KAN_input__s42", "MLP_input__s42", "Recall@10"),
             ("Dunnhumby", "dunnhumby", "Cell", "KAN_GRU_cell__s42", "MLP_GRU_cell__s42", "Recall@10")]
    rows = [["Dataset", "Placement", "Metric", "KAN", "Matched MLP"]]
    for name, ds, place, k, mm, met in pairs:
        rows.append([name, place, met, m(ds, k, met), m(ds, mm, met)])
    w.table(rows, [1.3, 1.5, 1.2, 1.1, 1.3], "KAN placements against parameter-matched MLPs (test, seed 42).",
            align_right=(3, 4))
    w.heading("Appendix C: Issues List")
    w.table([["#", "Issue", "Status"],
             ["1", "Third seed (44) for the four key YooChoose models; seed 43 for the matched MLP-GRU cell", "Open"],
             ["2", "Full Instacart and yoochoose1/4 runs on the UPES GPU", "Open"],
             ["3", "Instacart clustering gives 12,058 groups vs 9,407 in the original paper", "Documented"],
             ["4", "Benchmark of torch.compile and float32 optimisations on the GPU", "Open"],
             ["5", "Results, conclusion and abstract chapters of the report", "In progress"],
             ["6", "Publication of the KAN-GRU cell and KAN-NBR results", "Planned"]],
            [0.4, 4.6, 1.4])


def build(page_numbers):
    doc = Document(TEMPLATE)
    fill_front_matter(doc, page_numbers)
    anchor = strip_instructions(doc)
    w = Writer(doc, anchor)
    write_content(w)
    doc.core_properties.title = "SRS: " + TITLE
    doc.core_properties.author = "Abhishek Yadav, Saksham Agrawal, Yash Dagar"
    doc.save(OUT)


def page_map():
    from pypdf import PdfReader
    with tempfile.TemporaryDirectory() as d:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", d, str(OUT)],
                       check=True, capture_output=True)
        reader = PdfReader(str(Path(d) / (OUT.stem + ".pdf")))
        pages = [p.extract_text() or "" for p in reader.pages]
    found = {}
    targets = {
        "Table of Content": "TABLE OF CONTENTS", "Revision History": "REVISION HISTORY",
        "1 Introduction": "1 INTRODUCTION", "1.1 Purpose of the Project": "1.1 Purpose of the Project",
        "1.2 Target Beneficiary": "1.2 Target Beneficiary", "1.3 Project Scope": "1.3 Project Scope",
        "1.4 References": "1.4 References", "2 Project Description": "2 PROJECT DESCRIPTION",
        "2.1 Reference Algorithm": "2.1 Reference Algorithm", "2.2 Data/ Data structure": "2.2 Characteristic of Data",
        "2.3 SWOT Analysis": "2.3 SWOT Analysis", "2.4 Project Features": "2.4 Project Features",
        "2.5 User Classes and Characteristics": "2.5 User Classes", "2.6 Design and Implementation Constraints":
        "2.6 Design and Implementation", "2.7 Design diagrams": "2.7 Design Diagrams",
        "2.8 Assumption and Dependencies": "2.8 Assumption", "3 System Requirements": "3 SYSTEM REQUIREMENTS",
        "3.1 User Interface": "3.1 User Interface", "3.2 Software Interface": "3.2 Software Interface",
        "3.3 Database Interface": "3.3 Database Interface", "3.4 Protocols": "3.4 Protocols",
        "4 Non-functional Requirements": "4 NON-FUNCTIONAL", "4.1 Performance requirements": "4.1 Performance",
        "4.2 Security requirements": "4.2 Security", "4.3 Software Quality Attributes": "4.3 Software Quality",
        "5 Other Requirements": "5 OTHER REQUIREMENTS", "Appendix A: Glossary": "APPENDIX A",
        "Appendix B: Analysis Model": "APPENDIX B", "Appendix C: Issues List": "APPENDIX C"}
    for label, needle in targets.items():
        first = 0 if label in ("Table of Content", "Revision History") else 3
        for i, txt in enumerate(pages[first:], start=first):
            if needle.lower() in " ".join(txt.split()).lower():
                found[_norm(label)] = i + 1
                break
    return found


if __name__ == "__main__":
    build({})
    nums = page_map()
    build(nums)
    print("wrote", OUT, "with", len(nums), "page numbers")
